"""
Calculadora de valorización (docs/adr/0009): fórmula, rutas públicas, exportación a PDF y planilla, plantillas por cuenta,
hoja de presupuesto extendida y contador anónimo de uso.
"""
from io import BytesIO
from unittest import mock

from django.core.cache import cache
from django.core.management import call_command
from django.test import SimpleTestCase, override_settings
from django.utils import timezone
from datetime import timedelta
from openpyxl import load_workbook
from rest_framework.test import APITestCase

from . import calculadora, comisiones, derechos_titular, impacto
from .api import jwt_utils, views_calculadora
from .models import Contratacion, ItemPresupuesto, PlantillaValorizacion, Publicaciones, UsoCalculadora
from .tests import _crear_region_comuna_tipo, _crear_usuario

EJEMPLO = {
    'materiales': [
        {'descripcion': 'Cañería', 'cantidad': 3, 'unidad': 'm', 'precio_unitario': 2500, 'reutilizado': False},
        {'descripcion': 'Llave', 'cantidad': 1, 'precio_unitario': 8000, 'reutilizado': True},
    ],
    'traslado': {'modo': 'km', 'km_ida_vuelta': 20, 'costo_km': 250, 'peajes_estacionamiento': 1000},
    'mano_obra': {'horas': 3, 'valor_hora': 8000, 'personas': 1},
    'herramientas': {'modo': 'porcentaje', 'porcentaje': 5},
    'documento': 'BOLETA_HONORARIOS',
}


class FormulaCalculadoraTests(SimpleTestCase):
    def test_desglose_con_boleta_de_honorarios(self):
        r = calculadora.calcular(EJEMPLO)
        self.assertEqual(r['materiales_total'], 15500)          # 3 × 2.500 + 8.000
        self.assertEqual(r['traslado_total'], 6000)             # 20 km × 250 + 1.000
        self.assertEqual(r['mano_obra_total'], 24000)           # 3 h × 8.000
        self.assertEqual(r['herramientas_total'], 1200)         # 5 % de la mano de obra
        self.assertEqual(r['costo_directo'], 46700)
        self.assertEqual(r['imprevistos_monto'], 4670)          # 10 %
        self.assertEqual(r['ganancia_monto'], 7706)             # 15 % sobre (C + imprevistos), medio peso hacia arriba
        self.assertEqual(r['precio_neto'], 59076)
        self.assertEqual(r['precio_boleta'], 70329)             # 59.076 / (1 − 0,16) redondeado hacia arriba
        self.assertGreaterEqual(r['liquido'], r['precio_neto'])  # tras la retención quedan al menos los $59.076
        self.assertTrue(r['tiene_reutilizado'])

    def test_sin_documento_no_hay_retencion(self):
        r = calculadora.calcular({**EJEMPLO, 'documento': 'SIN_DOCUMENTO'})
        self.assertEqual(r['precio_boleta'], r['precio_neto'])
        self.assertEqual((r['retencion_monto'], r['retencion_pct']), (0, 0.0))
        self.assertEqual(r['liquido'], r['precio_neto'])

    def test_la_comision_es_la_misma_que_se_cobra_de_verdad(self):
        r = calculadora.calcular(EJEMPLO)
        desglose = comisiones.calcular(r['precio_boleta'])
        self.assertEqual(r['comision_monto'], desglose.comision)
        self.assertEqual(r['total_cliente'], r['precio_boleta'] + desglose.comision)

    def test_la_retencion_se_puede_cambiar_por_configuracion(self):
        with override_settings(CALCULADORA_RETENCION_BOLETA='10'):
            r = calculadora.calcular(EJEMPLO)
        self.assertEqual(r['retencion_pct'], 10.0)
        self.assertEqual(r['precio_boleta'], 65640)             # ceil(59.076 / 0,9)

    def test_valores_por_defecto_cuando_se_deja_vacio(self):
        r = calculadora.calcular({'mano_obra': {'horas': 2, 'valor_hora': 10000}})
        self.assertEqual((r['imprevistos_pct'], r['ganancia_pct']), (10.0, 15.0))
        self.assertEqual(r['precio_neto'], 25300)               # 20.000 × 1,10 × 1,15

    def test_traslado_fijo_y_herramientas_por_monto(self):
        r = calculadora.calcular({
            'traslado': {'modo': 'fijo', 'monto_fijo': 5000, 'km_ida_vuelta': 999, 'costo_km': 999},
            'herramientas': {'modo': 'monto', 'monto': 3000, 'porcentaje': 50},
            'imprevistos_pct': 0, 'ganancia_pct': 0,
        })
        self.assertEqual((r['traslado_total'], r['herramientas_total'], r['precio_neto']), (5000, 3000, 8000))

    def test_valores_en_cero_dan_cero_sin_comision(self):
        r = calculadora.calcular({})
        self.assertEqual((r['precio_boleta'], r['comision_monto'], r['total_cliente']), (0, 0, 0))

    def test_los_items_suman_exactamente_el_precio_a_cobrar(self):
        casos = [
            EJEMPLO,
            {**EJEMPLO, 'documento': 'SIN_DOCUMENTO'},
            {'materiales': [{'descripcion': 'Tornillos', 'cantidad': 7.5, 'precio_unitario': 133}], 'mano_obra': {'horas': 1.5, 'valor_hora': 7777, 'personas': 2},
             'imprevistos_pct': 7, 'ganancia_pct': 12.5, 'documento': 'BOLETA_HONORARIOS', 'retencion_pct': 13.75},
        ]
        for caso in casos:
            r = calculadora.calcular(caso)
            self.assertEqual(sum(i['monto'] for i in r['items']), r['precio_boleta'], caso)
            self.assertTrue(all(i['categoria'] in dict(ItemPresupuesto.CATEGORIAS_ITEM) for i in r['items']))

    def test_rechaza_datos_invalidos(self):
        malos = [
            {'mano_obra': {'horas': -1}},
            {'mano_obra': {'valor_hora': 'abc'}},
            {'materiales': [{'descripcion': 'x', 'precio_unitario': 10 ** 12}]},
            {'documento': 'FACTURA'},
            {'materiales': 'no es una lista'},
            {'imprevistos_pct': 101},
            'no es un objeto',
        ]
        for datos in malos:
            with self.assertRaises(calculadora.DatosInvalidos, msg=datos):
                calculadora.calcular(datos)

    def test_acepta_coma_decimal_y_filas_vacias_se_ignoran(self):
        r = calculadora.calcular({'materiales': [{'descripcion': '', 'precio_unitario': ''}, {'descripcion': 'Cable', 'cantidad': '2,5', 'precio_unitario': '1000'}]})
        self.assertEqual(r['materiales_total'], 2500)
        self.assertEqual(len(r['materiales']), 1)


class RutasPublicasCalculadoraTests(APITestCase):
    def tearDown(self):
        cache.clear()

    def test_config_es_publica(self):
        resp = self.client.get('/api/calculadora/config/')
        self.assertEqual(resp.status_code, 200)
        self.assertEqual(resp.data['retencion_porcentaje'], 16.0)
        self.assertEqual(resp.data['comision_tramos'][0], [50000, 6.0])

    def test_calcular_no_pide_cuenta(self):
        resp = self.client.post('/api/calculadora/calcular/', EJEMPLO, format='json')
        self.assertEqual(resp.status_code, 200)
        self.assertEqual(resp.data['precio_boleta'], 70329)

    def test_calcular_con_datos_malos_devuelve_400_con_mensaje(self):
        resp = self.client.post('/api/calculadora/calcular/', {'mano_obra': {'horas': -3}}, format='json')
        self.assertEqual(resp.status_code, 400)
        self.assertIn('horas', resp.data['detail'])

    def test_limite_de_uso_por_ip(self):
        with mock.patch.object(views_calculadora, 'LIMITE_CALCULOS', 3):
            codigos = [self.client.post('/api/calculadora/calcular/', {}, format='json').status_code for _ in range(4)]
        self.assertEqual(codigos, [200, 200, 200, 429])


class ExportacionCalculadoraTests(APITestCase):
    def tearDown(self):
        cache.clear()

    def test_xlsx_trae_formulas_visibles_y_los_datos_ingresados(self):
        resp = self.client.post('/api/calculadora/exportar/xlsx/', EJEMPLO, format='json')
        self.assertEqual(resp.status_code, 200)
        self.assertIn('attachment', resp['Content-Disposition'])
        hoja = load_workbook(BytesIO(resp.content)).active
        celdas = {c.coordinate: c.value for fila in hoja.iter_rows() for c in fila if c.value is not None}
        formulas = [v for v in celdas.values() if isinstance(v, str) and v.startswith('=')]
        self.assertGreaterEqual(len(formulas), 15)
        self.assertIn('Cañería', celdas.values())
        self.assertIn(2500, celdas.values())
        self.assertTrue(any('ROUNDUP' in f for f in formulas))      # precio con boleta
        self.assertTrue(any('MAX(1000' in f for f in formulas))     # comisión mínima de comisiones.py
        self.assertTrue(any('IF(' in f and '50000' in f for f in formulas))  # tramos de comisión

    def test_pdf_es_un_pdf(self):
        resp = self.client.post('/api/calculadora/exportar/pdf/', EJEMPLO, format='json')
        self.assertEqual(resp.status_code, 200)
        self.assertTrue(resp.content.startswith(b'%PDF'))

    def test_exportar_datos_invalidos_devuelve_400(self):
        self.assertEqual(self.client.post('/api/calculadora/exportar/pdf/', {'documento': 'X'}, format='json').status_code, 400)
        self.assertEqual(self.client.post('/api/calculadora/exportar/xlsx/', {'mano_obra': {'horas': 'x'}}, format='json').status_code, 400)

    def test_el_pdf_soporta_tildes_y_caracteres_raros(self):
        datos = {**EJEMPLO, 'materiales': [{'descripcion': 'Tubería “especial” – 50 mm ñandú', 'cantidad': 1, 'precio_unitario': 1000}]}
        self.assertEqual(self.client.post('/api/calculadora/exportar/pdf/', datos, format='json').status_code, 200)


class ContadorAnonimoTests(APITestCase):
    def tearDown(self):
        cache.clear()

    def _calcular(self, visitante=None):
        cuerpo = {**EJEMPLO}
        if visitante:
            cuerpo['visitante'] = visitante
        return self.client.post('/api/calculadora/calcular/', cuerpo, format='json')

    def _exportar(self, visitante, **extra):
        return self.client.post('/api/calculadora/exportar/pdf/', {**EJEMPLO, 'visitante': visitante, **extra}, format='json')

    def test_calcular_registra_la_visita_pero_no_una_valorizacion(self):
        self._calcular('visitante-0001')
        self._calcular('visitante-0001')
        self.assertEqual(UsoCalculadora.objects.count(), 1)
        self.assertEqual(UsoCalculadora.objects.get().valorizaciones, 0)

    def test_exportar_cuenta_una_valorizacion_con_boleta_y_reutilizado(self):
        self._exportar('visitante-0001')
        self._exportar('visitante-0001', documento='SIN_DOCUMENTO')
        uso = UsoCalculadora.objects.get()
        self.assertEqual((uso.valorizaciones, uso.con_boleta, uso.con_reutilizado), (2, 1, 2))

    def test_visitantes_distintos_son_filas_distintas(self):
        for i in range(3):
            self._calcular(f'visitante-000{i}')
        self.assertEqual(UsoCalculadora.objects.count(), 3)

    def test_no_se_guarda_ningun_dato_personal(self):
        self._calcular('visitante-secreto-123')
        uso = UsoCalculadora.objects.get()
        self.assertNotIn('secreto', uso.visitante)
        self.assertEqual(len(uso.visitante), 64)                      # hash
        campos = {f.name for f in UsoCalculadora._meta.get_fields()}
        self.assertFalse(campos & {'ip', 'usuario', 'email', 'datos'})

    def test_el_identificador_cambia_cada_mes(self):
        with mock.patch.object(views_calculadora.timezone, 'now', return_value=timezone.now().replace(year=2026, month=1, day=10)):
            self._calcular('visitante-0001')
        with mock.patch.object(views_calculadora.timezone, 'now', return_value=timezone.now().replace(year=2026, month=2, day=10)):
            self._calcular('visitante-0001')
        hashes = list(UsoCalculadora.objects.values_list('visitante', flat=True))
        self.assertEqual(len(set(hashes)), 2)

    def test_identificador_invalido_no_rompe_y_cae_en_un_solo_visitante(self):
        self._calcular('x')
        self._calcular('@@ mal @@')
        self._calcular()
        self.assertEqual(UsoCalculadora.objects.count(), 1)

    def test_el_reporte_suprime_menos_de_cinco_visitantes(self):
        for i in range(3):
            self._exportar(f'visitante-000{i}')
        calc = impacto.calcular()['calculadora']
        self.assertIsNone(list(calc['visitantes_por_mes'].values())[0])
        self.assertIsNone(calc['valorizaciones'])

    def test_el_reporte_publica_con_cinco_o_mas(self):
        for i in range(5):
            self._exportar(f'visitante-000{i}')
        calc = impacto.calcular()['calculadora']
        self.assertEqual(list(calc['visitantes_por_mes'].values())[0], 5)
        self.assertEqual((calc['valorizaciones'], calc['valorizaciones_con_boleta'], calc['valorizaciones_con_material_reutilizado']), (5, 5, 5))

    def test_la_purga_borra_el_contador_viejo(self):
        self._calcular('visitante-0001')
        UsoCalculadora.objects.update(actualizado_en=timezone.now() - timedelta(days=500))
        call_command('purgar_datos_vencidos')
        self.assertEqual(UsoCalculadora.objects.count(), 0)


class PlantillasValorizacionTests(APITestCase):
    def setUp(self):
        _, self.comuna, self.tipo = _crear_region_comuna_tipo()
        self.usuario = _crear_usuario('plantillas@test.com', es_proveedor=True, comuna=self.comuna, tipo_cuenta=self.tipo)
        self.otro = _crear_usuario('otro_pl@test.com', es_proveedor=True, comuna=self.comuna, tipo_cuenta=self.tipo)

    def tearDown(self):
        cache.clear()

    def _como(self, usuario):
        self.client.credentials(HTTP_AUTHORIZATION=f'Bearer {jwt_utils.generar_access_token(usuario)}')

    def test_requiere_cuenta(self):
        self.assertEqual(self.client.get('/api/calculadora/plantillas/').status_code, 401)
        self.assertEqual(self.client.post('/api/calculadora/plantillas/', {}, format='json').status_code, 401)

    def test_guardar_listar_y_reemplazar_por_nombre(self):
        self._como(self.usuario)
        resp = self.client.post('/api/calculadora/plantillas/', {'nombre': 'Gasfitería básica', 'tipo_trabajo': 'Gasfitería', 'datos': EJEMPLO}, format='json')
        self.assertEqual(resp.status_code, 201)
        resp = self.client.post('/api/calculadora/plantillas/', {'nombre': 'Gasfitería básica', 'datos': {**EJEMPLO, 'documento': 'SIN_DOCUMENTO'}}, format='json')
        self.assertEqual(resp.status_code, 200)
        lista = self.client.get('/api/calculadora/plantillas/').data
        self.assertEqual(len(lista), 1)
        self.assertEqual(lista[0]['datos']['documento'], 'SIN_DOCUMENTO')

    def test_guardar_exige_nombre_y_datos_validos(self):
        self._como(self.usuario)
        self.assertEqual(self.client.post('/api/calculadora/plantillas/', {'nombre': '  ', 'datos': EJEMPLO}, format='json').status_code, 400)
        self.assertEqual(self.client.post('/api/calculadora/plantillas/', {'nombre': 'x', 'datos': {'documento': 'Z'}}, format='json').status_code, 400)

    def test_cada_cuenta_ve_solo_las_suyas_y_ajenas_dan_404(self):
        ajena = PlantillaValorizacion.objects.create(usuario=self.otro, nombre='Ajena', datos=EJEMPLO)
        self._como(self.usuario)
        self.assertEqual(self.client.get('/api/calculadora/plantillas/').data, [])
        self.assertEqual(self.client.delete(f'/api/calculadora/plantillas/{ajena.id_plantilla}/').status_code, 404)
        self.assertEqual(self.client.get(f'/api/calculadora/plantillas/{ajena.id_plantilla}/items/').status_code, 404)
        self.assertTrue(PlantillaValorizacion.objects.filter(pk=ajena.pk).exists())

    def test_borrar_una_propia(self):
        propia = PlantillaValorizacion.objects.create(usuario=self.usuario, nombre='Mía', datos=EJEMPLO)
        self._como(self.usuario)
        self.assertEqual(self.client.delete(f'/api/calculadora/plantillas/{propia.id_plantilla}/').status_code, 204)
        self.assertFalse(PlantillaValorizacion.objects.exists())

    def test_tope_de_plantillas_por_cuenta(self):
        for i in range(views_calculadora.MAX_PLANTILLAS):
            PlantillaValorizacion.objects.create(usuario=self.usuario, nombre=f'P{i}', datos={})
        self._como(self.usuario)
        self.assertEqual(self.client.post('/api/calculadora/plantillas/', {'nombre': 'Una más', 'datos': EJEMPLO}, format='json').status_code, 400)
        # reemplazar una existente sí se puede
        self.assertEqual(self.client.post('/api/calculadora/plantillas/', {'nombre': 'P0', 'datos': EJEMPLO}, format='json').status_code, 200)

    def test_items_de_la_plantilla_para_la_hoja_de_presupuesto(self):
        propia = PlantillaValorizacion.objects.create(usuario=self.usuario, nombre='Mía', datos=EJEMPLO)
        self._como(self.usuario)
        resp = self.client.get(f'/api/calculadora/plantillas/{propia.id_plantilla}/items/')
        self.assertEqual(resp.status_code, 200)
        self.assertEqual(resp.data['total'], 70329)
        self.assertEqual(sum(i['monto'] for i in resp.data['items']), 70329)

    def test_aparecen_en_mis_datos_y_se_borran_al_anonimizar(self):
        PlantillaValorizacion.objects.create(usuario=self.usuario, nombre='Mía', datos=EJEMPLO)
        self.assertEqual(derechos_titular.exportar_datos(self.usuario)['plantillas_de_valorizacion'][0]['nombre'], 'Mía')
        derechos_titular.eliminar_cuenta(self.usuario)
        self.assertFalse(PlantillaValorizacion.objects.filter(usuario=self.usuario).exists())


class HojaDePresupuestoExtendidaTests(APITestCase):
    def setUp(self):
        _, self.comuna, self.tipo = _crear_region_comuna_tipo()
        self.proveedor = _crear_usuario('prov_hoja@test.com', 'clave_prov', es_proveedor=True, comuna=self.comuna, tipo_cuenta=self.tipo)
        self.cliente = _crear_usuario('cli_hoja@test.com', comuna=self.comuna, tipo_cuenta=self.tipo)
        pub = Publicaciones.objects.create(usuario_publicador=self.proveedor, titulo='Gasfitería', estado_moderacion=Publicaciones.APROBADA, precio=15000)
        self.contratacion = Contratacion.objects.create(publicacion=pub, cliente=self.cliente, proveedor=self.proveedor, estado=Contratacion.SOLICITADA)
        self.client.credentials(HTTP_AUTHORIZATION=f'Bearer {jwt_utils.generar_access_token(self.proveedor)}')

    def tearDown(self):
        cache.clear()

    def test_confirmar_con_los_items_de_la_calculadora_guarda_el_detalle(self):
        items = calculadora.calcular(EJEMPLO)['items']
        resp = self.client.post(f'/api/contrataciones/{self.contratacion.id_contratacion}/confirmar/', {'password': 'clave_prov', 'items': items}, format='json')
        self.assertEqual(resp.status_code, 200, resp.data)
        self.contratacion.refresh_from_db()
        self.assertEqual(self.contratacion.monto_acordado, 70329)
        guardados = {i.descripcion: i for i in self.contratacion.items_presupuesto.all()}
        canería = guardados['Cañería']
        self.assertEqual((float(canería.cantidad), canería.unidad, canería.precio_unitario, canería.reutilizado), (3.0, 'm', 2500, False))
        self.assertTrue(guardados['Llave'].reutilizado)
        self.assertEqual(guardados['Herramientas'].categoria, ItemPresupuesto.HERRAMIENTAS)
        self.assertEqual(guardados['Imprevistos (10 %)'].categoria, ItemPresupuesto.IMPREVISTOS)
        detalle = self.client.get(f'/api/contrataciones/{self.contratacion.id_contratacion}/').data
        self.assertEqual(detalle['items_presupuesto'][0]['unidad'], 'm')

    def test_las_filas_antiguas_sin_detalle_siguen_funcionando(self):
        resp = self.client.post(f'/api/contrataciones/{self.contratacion.id_contratacion}/confirmar/', {
            'password': 'clave_prov', 'items': [{'descripcion': 'Mano de obra', 'categoria': 'MANO_DE_OBRA', 'monto': 12000}],
        }, format='json')
        self.assertEqual(resp.status_code, 200)
        item = self.contratacion.items_presupuesto.get()
        self.assertEqual((item.cantidad, item.unidad, item.precio_unitario, item.reutilizado), (None, '', None, False))

    def test_detalle_basura_se_ignora_sin_romper(self):
        resp = self.client.post(f'/api/contrataciones/{self.contratacion.id_contratacion}/confirmar/', {
            'password': 'clave_prov',
            'items': [{'descripcion': 'Algo', 'monto': 5000, 'cantidad': 'mucho', 'precio_unitario': 'caro', 'reutilizado': 'sí'}],
        }, format='json')
        self.assertEqual(resp.status_code, 200)
        item = self.contratacion.items_presupuesto.get()
        self.assertEqual((item.cantidad, item.precio_unitario), (None, None))


class ModoNormalCalculadoraTests(SimpleTestCase):
    """Modo normal: la persona ya sabe cuánto cobrar; solo se ve el efecto de la boleta y la comisión."""

    def test_sin_documento_el_precio_es_el_que_se_escribe(self):
        r = calculadora.calcular({'modo': 'simple', 'precio_simple': 30000})
        self.assertTrue(r['simple'])
        self.assertEqual((r['precio_neto'], r['precio_boleta'], r['liquido']), (30000, 30000, 30000))
        self.assertEqual((r['imprevistos_monto'], r['ganancia_monto']), (0, 0))
        self.assertEqual(r['total_cliente'], 30000 + comisiones.calcular(30000).comision)
        self.assertEqual(r['items'], [{'descripcion': 'Servicio', 'categoria': 'OTRO', 'monto': 30000}])

    def test_con_boleta_sube_el_precio_para_que_quede_lo_deseado(self):
        r = calculadora.calcular({'modo': 'simple', 'precio_simple': 30000, 'documento': 'BOLETA_HONORARIOS'})
        self.assertEqual(r['precio_neto'], 30000)
        self.assertEqual(r['precio_boleta'], 35715)             # ceil(30.000 / 0,84)
        self.assertGreaterEqual(r['liquido'], 30000)
        self.assertEqual(sum(i['monto'] for i in r['items']), r['precio_boleta'])

    def test_ignora_lo_de_la_calculadora_de_costos(self):
        r = calculadora.calcular({**EJEMPLO, 'modo': 'simple', 'precio_simple': 10000, 'imprevistos_pct': 50, 'ganancia_pct': 50})
        self.assertEqual(r['precio_neto'], 10000)
        self.assertEqual(r['materiales'], [])

    def test_exige_un_numero_valido(self):
        with self.assertRaises(calculadora.DatosInvalidos):
            calculadora.calcular({'modo': 'simple', 'precio_simple': -5})
        self.assertEqual(calculadora.calcular({'modo': 'simple'})['precio_boleta'], 0)

    def test_el_modo_de_costos_sigue_igual(self):
        self.assertFalse(calculadora.calcular(EJEMPLO)['simple'])
        self.assertEqual(calculadora.calcular(EJEMPLO)['precio_boleta'], 70329)


class PrecioDeLaPublicacionEnLaContratacionTests(APITestCase):
    """La calculadora de un trabajo parte del precio de la publicación: tiene que viajar con la contratación."""

    def test_la_contratacion_trae_el_precio_de_la_publicacion(self):
        _, comuna, tipo = _crear_region_comuna_tipo()
        proveedor = _crear_usuario('prov_precio@test.com', es_proveedor=True, comuna=comuna, tipo_cuenta=tipo)
        cliente = _crear_usuario('cli_precio@test.com', comuna=comuna, tipo_cuenta=tipo)
        con_precio = Publicaciones.objects.create(usuario_publicador=proveedor, titulo='A', estado_moderacion=Publicaciones.APROBADA, precio=25000)
        sin_precio = Publicaciones.objects.create(usuario_publicador=proveedor, titulo='B', estado_moderacion=Publicaciones.APROBADA)
        c1 = Contratacion.objects.create(publicacion=con_precio, cliente=cliente, proveedor=proveedor)
        c2 = Contratacion.objects.create(publicacion=sin_precio, cliente=cliente, proveedor=proveedor)
        self.client.credentials(HTTP_AUTHORIZATION=f'Bearer {jwt_utils.generar_access_token(proveedor)}')
        self.assertEqual(self.client.get(f'/api/contrataciones/{c1.id_contratacion}/').data['publicacion_precio'], 25000)
        self.assertIsNone(self.client.get(f'/api/contrataciones/{c2.id_contratacion}/').data['publicacion_precio'])
        self.assertEqual({c['id_contratacion']: c['publicacion_precio'] for c in self.client.get('/api/contrataciones/').data}[c1.id_contratacion], 25000)


class IvaYDocumentosCalculadoraTests(SimpleTestCase):
    """La boleta de honorarios no lleva IVA; un documento afecto sí (19 %) y ese IVA no es de quien cobra."""

    BASE = {'mano_obra': {'horas': 3, 'valor_hora': 10000}, 'imprevistos_pct': 0, 'ganancia_pct': 0}

    def test_con_iva_se_suma_encima_y_no_cuenta_en_lo_que_recibes(self):
        r = calculadora.calcular({**self.BASE, 'documento': 'DOCUMENTO_IVA'})
        self.assertEqual(r['precio_neto'], 30000)
        self.assertEqual((r['iva_pct'], r['iva_monto']), (19.0, 5700))
        self.assertEqual(r['precio_boleta'], 35700)
        self.assertEqual((r['retencion_monto'], r['liquido']), (0, 30000))
        self.assertEqual(r['total_cliente'], 35700 + comisiones.calcular(35700).comision)  # la comisión se calcula sobre lo que se cobra

    def test_la_boleta_de_honorarios_no_lleva_iva(self):
        r = calculadora.calcular({**self.BASE, 'documento': 'BOLETA_HONORARIOS'})
        self.assertEqual((r['iva_pct'], r['iva_monto']), (0.0, 0))
        self.assertGreater(r['retencion_monto'], 0)

    def test_sin_documento_no_hay_ni_iva_ni_retencion(self):
        r = calculadora.calcular({**self.BASE, 'documento': 'SIN_DOCUMENTO'})
        self.assertEqual((r['iva_monto'], r['retencion_monto'], r['precio_boleta']), (0, 0, 30000))

    def test_los_items_suman_el_precio_con_iva(self):
        r = calculadora.calcular({**EJEMPLO, 'documento': 'DOCUMENTO_IVA'})
        self.assertEqual(sum(i['monto'] for i in r['items']), r['precio_boleta'])
        self.assertTrue(any(i['descripcion'].startswith('IVA') for i in r['items']))

    def test_imprevistos_y_ganancia_se_pueden_dejar_en_cero(self):
        r = calculadora.calcular({**self.BASE})
        self.assertEqual((r['imprevistos_monto'], r['ganancia_monto'], r['precio_neto']), (0, 0, 30000))
        r = calculadora.calcular({'mano_obra': {'horas': 3, 'valor_hora': 10000}})  # sin decirlo: por defecto 10 % y 15 %
        self.assertEqual(r['precio_neto'], 37950)

    def test_modo_normal_con_iva(self):
        r = calculadora.calcular({'modo': 'simple', 'precio_simple': 20000, 'documento': 'DOCUMENTO_IVA'})
        self.assertEqual((r['precio_boleta'], r['liquido']), (23800, 20000))

    def test_el_iva_se_configura(self):
        with override_settings(CALCULADORA_IVA='10'):
            r = calculadora.calcular({**self.BASE, 'documento': 'DOCUMENTO_IVA'})
        self.assertEqual(r['iva_monto'], 3000)
